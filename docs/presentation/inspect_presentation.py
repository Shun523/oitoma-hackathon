"""Check the PPTX and its real LibreOffice PDF render; write page previews."""

import json
from pathlib import Path
import re
import zipfile

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


ROOT = Path(__file__).resolve().parent
PREVIEWS = ROOT/'previews'
PREVIEWS.mkdir(exist_ok=True)
PPTX = ROOT/'oitoma-4min.pptx'
PDF = ROOT/'oitoma-4min.pdf'
data = json.loads((ROOT/'slide-data.json').read_text())
preflight = json.loads((ROOT/'preflight.json').read_text())
prs = Presentation(PPTX)
pdf = fitz.open(PDF)
issues = list(preflight['issues'])
if len(prs.slides) != 10 or len(pdf) != 10:
    issues.append('Expected 10 PPTX slides and 10 PDF pages.')


def normalize(text):
    return re.sub(r'\s', '', text)


notes_ok = True
pages = []
for i, (slide, page) in enumerate(zip(prs.slides, pdf)):
    expected_notes = ['【読み上げ原稿】', '【根拠】', '【想定質問と回答】', '【次へのつなぎ】']
    notes = slide.notes_slide.notes_text_frame.text
    for expected in expected_notes:
        if expected not in notes:
            notes_ok = False
            issues.append(f'Slide {i+1}: missing notes section {expected}')
    page_text = normalize(page.get_text())
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            expected = normalize(paragraph.text)
            if expected and expected not in page_text:
                issues.append(f'Slide {i+1}: PDF is missing text: {paragraph.text}')
    blocks = page.get_text('dict')['blocks']
    fonts = set()
    for block in blocks:
        for line in block.get('lines', []):
            for span in line.get('spans', []):
                fonts.add(span['font'])
                bbox=fitz.Rect(span['bbox'])
                if bbox.x0 < -0.2 or bbox.y0 < -0.2 or bbox.x1 > page.rect.width+.2 or bbox.y1 > page.rect.height+.2:
                    issues.append(f'Slide {i+1}: PDF glyph outside canvas: {span["text"]}')
    path=PREVIEWS/f'slide-{i+1:02}.png'
    page.get_pixmap(matrix=fitz.Matrix(2,2), alpha=False).save(path)
    pages.append(dict(slide=i+1, width_pt=page.rect.width, height_pt=page.rect.height,
                      preview=str(path.relative_to(ROOT)), fonts=sorted(fonts)))

font = ImageFont.truetype('/usr/share/fonts/opentype/ipaexfont-gothic/ipaexg.ttf', 19)
cw,ch,gap,label=800,450,22,36
sheet=Image.new('RGB', (cw*2+gap*3, (ch+label)*5+gap*6), '#E8EDF0')
draw=ImageDraw.Draw(sheet)
for i in range(10):
    x=gap+(i%2)*(cw+gap)
    y=gap+(i//2)*(ch+label+gap)
    image=Image.open(PREVIEWS/f'slide-{i+1:02}.png').convert('RGB')
    image.thumbnail((cw,ch), Image.Resampling.LANCZOS)
    sheet.paste(image,(x,y))
    title=f'{i+1:02}  '+data[i]['title']
    draw.text((x,y+ch+6),title,font=font,fill='#172C38')
sheet.save(PREVIEWS/'contact-sheet.png')

with zipfile.ZipFile(PPTX) as zf:
    for member in zf.infolist():
        if member.filename.endswith('.xml'):
            from lxml import etree
            etree.fromstring(zf.read(member.filename))
    notes_parts=[n for n in zf.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)]
    if len(notes_parts)!=10:
        issues.append('Expected 10 embedded notes XML parts.')

result=dict(issues=issues, pages=pages, notes_complete=notes_ok,
            embedded_notes_count=len(notes_parts),
            total_main_seconds=preflight['main_seconds'],
            body_counts=preflight['body_counts'],
            minimum_content_font_pt=preflight['minimum_content_font_pt'],
            pptx_native_shapes=True)
(ROOT/'render-check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k!='pages'},ensure_ascii=False,indent=2))
if issues:
    raise SystemExit(1)
